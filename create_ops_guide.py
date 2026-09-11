from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "/Users/hakaur/media_planner/Media Planner Operations Guide.docx"

NAVY = "17365D"
BLUE = "DCEAF7"
PALE = "F5F8FC"
GRAY = "D9E1EA"
TEXT = "1F2937"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def borders(table, color=GRAY):
    tbl_pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), color)
        el.append(edge)
    tbl_pr.append(el)

def set_cell_margin(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def style_cell(cell, header=False, center=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    if header:
        shade(cell, NAVY)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Aptos"
            r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            r.font.size = Pt(8.7)
            r.font.color.rgb = RGBColor.from_string("FFFFFF" if header else TEXT)
            r.bold = header

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.text = header
        style_cell(cell, header=True, center=idx == 0)
    set_repeat_table_header(table.rows[0])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            cells[idx].text = value
            if row_idx % 2:
                shade(cells[idx], PALE)
            style_cell(cells[idx], center=idx == 0)
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.22 + 0.20 * level)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.add_run(text)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)
    return p

def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.60)
sec.left_margin = Inches(0.70)
sec.right_margin = Inches(0.70)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
    st = styles[name]
    st.font.name = "Aptos Display" if name == "Title" else "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True

# Title page
p = doc.add_paragraph(style="Title")
p.paragraph_format.space_after = Pt(5)
p.add_run("Media Planner Operations Guide")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("How the planner selects, allocates, and validates on deck media inventory")
r.italic = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string("4B5563")

add_para(doc, "This guide explains the rules currently used by the Media Planner so Operations can validate whether they reflect booking policy. The planner is designed to create a feasible on deck plan from campaign inputs, forecast availability, booking commitments, historical performance, and the rate card.")
add_para(doc, "", "What Operations should review. ")
add_para(doc, "Please use this document to confirm the hard limits, prioritisation logic, and operational constraints below. If a policy rule is missing or needs adjustment, share the desired rule, its business rationale, scope, and any exceptions so it can be added deliberately.")

add_heading(doc, "Planner input and output", 1)
add_table(doc,
    ["Input", "How it is used"],
    [
        ["Campaign details", "Brand, country, commercial category, flight dates, currency, budget, discount, objective, and optional phase splits define the planning scope."],
        ["Availability", "For each date and slot, forecast views are reduced by already booked views. The remaining views are the capacity available to plan."],
        ["Historical performance", "Historical delivery, clicks, spend, revenue, recency, brand context, and category relevance provide the performance signals used to rank eligible slots."],
        ["Slot metadata", "Page, category, zone, marketplace, and supported pricing models are used to ensure that selected slots fit the requested campaign."],
        ["Rate card", "CPM and CPD prices determine whether a candidate fits the requested dates and the remaining campaign budget."],
    ], [1.5, 5.5])

add_heading(doc, "Planning flow", 1)
for text in [
    "Validate the campaign inputs and create the campaign phase structure. If no phases are supplied, the full flight is treated as one phase.",
    "Load the eligible slot catalogue, historical performance, forecast inventory, current bookings, and price data for the selected countries and dates.",
    "Calculate remaining capacity by date and slot as forecast views minus booked views. A slot with no remaining forecast capacity cannot be placed.",
    "Score and order eligible slots according to the selected campaign objective, while retaining a usable mix of homepage, category, and other relevant placements.",
    "Allocate budget across country, brand, phase, marketplace, commercial category, and objective tracks. Build plan lines only when they satisfy capacity, pricing, date, and budget rules.",
    "Return the generated plan with allocation diagnostics and clear reasons for any selected slots that could not be placed.",
]:
    add_bullet(doc, text)

doc.add_page_break()

add_heading(doc, "Campaign objective rules", 1)
add_para(doc, "The campaign objective changes the order in which eligible placements are considered. This makes the recommendation responsive to the actual eligible inventory, price, and performance data for the flight.")
add_table(doc,
    ["Objective", "Primary priority", "Placement preference"],
    [
        ["Reach or visibility", "Higher visibility and higher reach signals", "Homepage placements receive the strongest preference. The ordering cycle gives homepage placements three opportunities before category and other placements."],
        ["RoAS or performance", "Historical return and commercial category relevance", "CLP, PLP, category, and sale-page placements receive the strongest preference. The ordering cycle gives category placements three opportunities before homepage and other placements."],
        ["CTR", "Historical click-through performance and category relevance", "Uses the same category-led placement preference as performance planning."],
        ["Balanced", "Combined performance signals", "Uses a mixed ordering cycle across homepage, category, and other placements."],
    ], [1.35, 2.5, 3.15])

add_heading(doc, "How budget size changes the plan", 1)
add_table(doc,
    ["Budget band", "Planner behaviour", "Operational expectation"],
    [
        ["Small budget\nUSD 10,000 or below", "The planner targets at least 6 on deck lines per country, subject to capacity and the hard limits below. It prioritises efficient, eligible slots and avoids concentrating all spend in one placement.", "Expect a focused plan. Premium CPD placements may not fit the budget; CPM and lower-cost eligible inventory are more likely to be used."],
        ["Large budget\nAbove USD 10,000", "The planner targets at least 10 on deck lines per country, subject to capacity and the hard limits below. It has more room to diversify across placement families and allocation buckets.", "Expect broader coverage across phase, marketplace, commercial category, and objective tracks where eligible inventory exists."],
    ], [1.4, 3.0, 2.6])
add_para(doc, "These line counts are diversity targets, not a promise. The planner will return fewer lines when there is insufficient eligible inventory, no valid rate, an unavoidable date overlap, or a budget constraint that prevents the minimum purchasable amount.")

add_heading(doc, "Allocation rules", 1)
for text in [
    "Countries share the on deck budget equally unless the campaign structure creates a more specific allocation bucket.",
    "When selected, brand, phase, marketplace, commercial category, and objective split inputs are applied as allocation targets. The planner reports actual versus requested splits so Operations can identify the closest feasible result.",
    "Generated slots prefer candidates not already used elsewhere in the campaign, after the normal allocation filters. Reuse remains available when needed to meet the allocation target and budget utilisation.",
    "Manual slot additions are exempt from this campaign-wide diversity preference.",
    "A user-selected slot is represented where possible, but it must still have an available supported pricing model, a valid date-specific rate, and capacity. Manual selections can be used to override normal recommendation eligibility.",
    "The same slot is not reused in the same phase. It can be used again only in a non-overlapping phase when the rest of the rules permit it.",
    "The plan preserves selected manual and locked lines during regeneration; other lines can be regenerated around them.",
]:
    add_bullet(doc, text)

add_heading(doc, "Hard limits and validation rules", 1)
add_para(doc, "The following controls are applied to protect feasibility and prevent a plan from overbooking inventory or concentrating the budget too heavily in one slot.")
add_table(doc,
    ["Rule", "Current setting", "What happens if it is not met"],
    [
        ["Campaign dates", "End date must not be before start date; start date must be today or later.", "The request is rejected before planning."],
        ["Required planning scope", "A positive on deck budget, at least one commercial category, and at least one supported country are required.", "The request is rejected before planning."],
        ["Supported countries", "AE, SA, and EG.", "Other country values are rejected."],
        ["Available inventory", "Only forecast views remaining after booked views are eligible. Minimum CPM block: 1,000 views.", "The slot is omitted if it has no capacity or cannot meet the 1,000-view minimum."],
        ["Per-slot spend cap", "Maximum 35% of the total on deck budget per slot.", "Additional spend on that slot is capped or rejected. FOC slots are exempt from this spend cap."],
        ["CPD suitability", "Automatically recommended CPD inventory requires an on deck budget of at least USD 15,000. Manual selection can override this recommendation threshold.", "Below the threshold, CPD is not automatically recommended; the planner favours eligible CPM options."],
        ["Pricing validity", "A valid CPM or CPD price must exist for the planned date window.", "The slot is omitted with a rate-related reason."],
        ["Duplicate placement control", "One slot cannot be used twice in the same phase. Another slot using the same country, marketplace, category, and zone is also blocked within that phase.", "The planner considers another eligible placement or reports the reason it could not place the slot."],
        ["Lines per phase", "Maximum 6 generated lines per phase allocation bucket.", "The planner stops adding generated lines in that bucket and uses the remaining eligible allocation paths."],
    ], [1.45, 2.85, 2.7])

add_heading(doc, "Date and flight handling", 1)
add_para(doc, "Campaign dates are handled as 09:00 to 09:00 service windows. For example, a flight from 09:00 on one date to 09:00 on the next is treated as one day. On partial boundary dates, forecast views, booked views, and CPD costs are prorated by the covered fraction of the day.")
add_para(doc, "For CPM lines, the planner calculates a purchasable volume within the available views and budget. For CPD lines, it adds date segments only while the date-specific daily cost fits the remaining allocation. This is why a CPD placement may run for fewer days than requested when the remaining budget cannot purchase the next day.")

add_heading(doc, "What Operations will see when a plan is constrained", 1)
for text in [
    "Budget utilisation, remaining amount, and actual allocation splits in the plan diagnostics.",
    "A closest feasible status when the requested phase, marketplace, or objective split cannot be matched exactly.",
    "Specific omission reasons such as no remaining forecast views, no valid rate, slot cap reached, date overlap, duplicate zone/category use, or insufficient budget for the first CPD segment or the minimum CPM block.",
    "A manually editable plan table after generation, allowing Operations to review slots, dates, pricing model, rate, views, spend, phase, and notes before using the plan operationally.",
]:
    add_bullet(doc, text)

add_heading(doc, "Operations review checklist", 1)
add_para(doc, "Please review the following questions and share any additions or changes as operational rules rather than one-off exceptions.")
questions = [
    "Are the supported countries, 35% per-slot cap, 1,000-view CPM minimum, and USD 15,000 automatic CPD threshold correct for current booking policy?",
    "Are there page, zone, marketplace, commercial-category, or brand exclusions that should always be enforced?",
    "Are there minimum campaign duration, minimum CPD booking duration, frequency, or impression-delivery rules that must be added?",
    "Should specific slot types have different concentration caps or budget thresholds?",
    "Are there rules for FOC inventory, premium events, or commercial commitments that should override the standard logic?",
    "What data fields and refresh cadence can Operations guarantee for the rate card, inventory forecast, and booked views?",
]
for q in questions:
    add_bullet(doc, q)

add_heading(doc, "Rate card and external data dependency", 1)
add_para(doc, "The planner depends on an accurate date-wise rate card for every slot. The rate card should specify the valid CPM and CPD rates by country, slot, and date range. This allows the planner to account for normal days, sale periods, and any date-specific commercial pricing through the price data itself, rather than through a separate sale-period rule engine.")
add_para(doc, "", "External dependency. ")
add_para(doc, "Operations and the relevant external teams must provide and maintain the rate-card table, including timely updates for sale windows, price changes, and unavailable inventory. If the planner does not receive a valid rate for the selected date window, it will not plan that slot. Forecast and booking data must also be refreshed so remaining capacity is accurate.")
add_para(doc, "Once this data is available, Operations can use the rule checklist above to validate the current behaviour and propose any additional permanent planning rules.")

# Footer with page field
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_run = footer.add_run("Media Planner Operations Guide  |  ")
footer_run.font.size = Pt(8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

doc.core_properties.title = "Media Planner Operations Guide"
doc.core_properties.subject = "Operations review of Media Planner rules and rate-card dependency"
doc.core_properties.author = "Media Planner"
doc.save(OUT)
print(OUT)
