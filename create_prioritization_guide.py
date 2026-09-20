from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Media Planner Rules Guide.docx"
NAVY = "17365D"
GRID = "D9D9D9"


def set_cell_border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID)
        borders.append(tag)
    props.append(borders)


def fill(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)


def write_cell(cell, text, bold=False, color=None):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def add_bullets(doc, items):
    for text in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(text)


def add_table(doc, headers, values, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, text in zip(table.rows[0].cells, headers):
        fill(cell, NAVY)
        write_cell(cell, text, bold=True, color="FFFFFF")
    for row in values:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            write_cell(cell, text)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(3)
for style_name, size in (("Title", 19), ("Heading 1", 13.5), ("Heading 2", 11)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(7)
styles["Heading 1"].paragraph_format.space_after = Pt(3)

title = doc.add_paragraph(style="Title")
title.add_run("Media Planner Selection and Budget Rules")
title_borders = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_borders.append(title_bottom)
title._p.get_or_add_pPr().append(title_borders)

subtitle = doc.add_paragraph()
subtitle_run = subtitle.add_run("Guide for Sales and Operations")
subtitle_run.italic = True
subtitle_run.font.size = Pt(11)
subtitle_run.font.color.rgb = RGBColor(89, 89, 89)
subtitle.paragraph_format.space_after = Pt(14)

doc.add_paragraph(
    "The planner chooses eligible slots first, then optimizes the best practical mix for the campaign. "
    "It uses the campaign objective, requested phase and marketplace splits, available inventory, and daily rates."
)

doc.add_heading("How the campaign objective guides slot choice", level=1)
add_table(
    doc,
    ("Campaign objective", "How slots are prioritized"),
    (
        ("Reach", "Gives stronger preference to Homepage and high-visibility placements."),
        ("ROAS or CTR", "Gives stronger preference to CLP and category-page placements."),
        ("Balanced objective", "Aims for a practical mix of Homepage, CLP, and other relevant placements."),
    ),
    (2.0, 4.8),
)
doc.add_heading("How budget is optimized", level=1)
doc.add_paragraph(
    "The planner aims to match the selected phase and Core or Supermall splits as closely as possible. These are optimization targets, not hard stops, so the planner can still create a workable plan when suitable inventory is limited."
)
add_bullets(doc, [
    "It aims to use the budget fully while still respecting inventory, pricing, and booking rules.",
    "It keeps a small continuity margin equal to the lower of 5 percent of the on-deck budget or USD 250. This gives the planner room to maintain sensible delivery across dates.",
    "If the exact phase or marketplace split is not feasible, the plan uses the closest achievable split and shows the difference.",
])

doc.add_heading("Recommendation count and manual selection", level=1)
add_table(
    doc,
    ("Campaign budget", "Starting recommendation"),
    (
        ("Up to USD 10,000", "6 recommended slots per selected country"),
        ("Above USD 10,000", "10 recommended slots per selected country"),
    ),
    (2.3, 4.5),
)
doc.add_paragraph(
    "These are starting recommendations, not hard slot counts. The planner can recommend additional eligible slots if the initial set does not use enough of the budget. Users can also add any slot manually."
)
doc.add_paragraph("Manual additions bypass recommendation filters, but still need valid inventory, campaign dates, and CPM or CPD pricing.")

doc.add_heading("CPD and Q4 pricing", level=1)
add_bullets(doc, [
    "Campaign service days run from 9 AM to 9 AM. For example, 10 September to 12 September is two chargeable days, while 18 September to 19 September is one chargeable day.",
    "Homepage CPD is recommended only when the campaign budget is USD 15,000 or above. This avoids spreading a premium Homepage placement too thinly across many days.",
    "CLP and category-page CPD can still be recommended below USD 15,000 when valid inventory and pricing are available.",
    "For October to December, Q4 daily rates are implemented. The planner reads the applicable daily CPM or CPD rate for the selected slot, country, and date.",
    "If a Q4 rate changes during a campaign, the plan shows separate rows for each date range with the same rate.",
])

doc.add_heading("Slot diversity", level=1)
doc.add_paragraph(
    "The planner prefers eligible slots that have not already been auto-selected elsewhere in the campaign. A previously used slot remains available as a fallback when it is the best feasible option. Manual selections are not deprioritized."
)

doc.add_heading("Why hard limits are not recommended", level=1)
doc.add_paragraph(
    "Making every phase split, marketplace split, daily-delivery target, or placement mix a hard limit can prevent the planner from creating a plan when the available inventory cannot satisfy every requirement. It can also leave budget unused or break delivery continuity. The planner therefore optimizes these targets and protects only the rules needed for a plan to be bookable."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Media Planner Selection and Budget Rules")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUTPUT)
print(OUTPUT)
