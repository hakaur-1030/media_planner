from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Media Planner Rules Guide.docx"
NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F6F8FB"
GRID = "D9E1EA"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_text(paragraph, text, bold=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    add_text(p, text)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, text, bold=True, color="000000", size=14)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text)
    return p


def make_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, bold=True, color="FFFFFF", size=9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            if row_index % 2:
                set_cell_shading(cell, PALE)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_text(p, value, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    add_text(title, "Media Planner Rules Guide", bold=True, color="000000", size=22)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    add_text(subtitle, "How slots are selected and budget is allocated", color="4A5568", size=11)

    add_body(
        doc,
        "This guide explains the rules used to create an on deck media plan. The planner first checks whether a slot can be booked, then ranks suitable options, and finally allocates the budget without exceeding inventory or budget limits.",
    )

    add_heading(doc, "How a slot is selected")
    add_body(doc, "The planner starts with the campaign brief: selected countries, ComCats, dates, phases, budget, marketplace split, and objective. It then keeps only slots that are relevant and can be delivered.")
    make_table(
        doc,
        ["Step", "Rule"],
        [
            ("Eligibility", "The slot must have forecast inventory in the selected phase, a valid CPM or CPD rate, and a matching country, marketplace, and ComCat."),
            ("Booking checks", "Recently unbooked or unavailable slots are removed. A user can still choose from the wider manual catalog when there is a commercial reason."),
            ("Performance", "The planner considers brand history where available, comparable ComCat performance, recent trends, and the confidence of the available data."),
            ("Objective", "Reach gives more weight to Homepage visibility. ROAS and CTR give more weight to CLP and category placements. A balanced campaign keeps a mix of placement types where inventory exists."),
            ("Variety", "The planner avoids repeatedly selecting the same automatically generated slot when other suitable options are available."),
        ],
        [1.15, 5.85],
    )

    add_heading(doc, "Number of recommended slots")
    add_body(doc, "The number below is a starting target for each selected country. It is not a maximum. The planner can add more eligible recommendations if the initial set cannot use at least 95 percent of the budget, and users can add manual slots at any time.")
    make_table(
        doc,
        ["On deck budget", "Starting recommendation target"],
        [
            ("Up to USD 10,000", "6 slots per country"),
            ("Above USD 10,000", "10 slots per country"),
        ],
        [2.8, 4.2],
    )

    add_heading(doc, "How the budget is split")
    add_body(doc, "The planner starts by dividing the on deck budget across selected countries. Within each country, it aims to follow the phase, marketplace, brand, ComCat, and objective splits entered in the brief. It then assigns spend to the strongest eligible slots within those buckets.")
    add_body(doc, "These are allocation targets, not guarantees. If inventory, slot eligibility, or a budget limit prevents an exact split, the planner uses the closest feasible result and reports the difference.")
    make_table(
        doc,
        ["Rule", "What it protects"],
        [
            ("Total budget", "The on deck plan cannot exceed the approved on deck budget."),
            ("Slot budget cap", "A normal paid slot cannot take more than 35 percent of the on deck budget."),
            ("Inventory", "Planned impressions cannot exceed the forecast inventory available in the assigned phase."),
            ("Placement conflicts", "The same slot cannot be used twice in the same phase. The same category and zone are also prevented from being duplicated within one country and phase."),
        ],
        [1.55, 5.45],
    )

    add_heading(doc, "CPM and CPD rules")
    add_bullet(doc, "CPM cost is calculated from planned impressions and the applicable CPM rate.")
    add_bullet(doc, "CPD cost is calculated from the applicable daily CPD rate and the booked days.")
    add_bullet(doc, "A CPD slot must have enough remaining budget for at least its first eligible day. A CPM slot must meet the minimum impression block.")
    add_bullet(doc, "Homepage CPD is recommended only for campaigns of at least USD 15,000. This avoids spreading premium Homepage delivery too thinly. CLP and category CPD can still be recommended below USD 15,000 when rate and inventory are available.")
    add_bullet(doc, "Users can manually choose Homepage CPD below USD 15,000 when there is a specific commercial requirement.")

    add_heading(doc, "Daily Q4 rate handling")
    add_body(doc, "From 1 October to 31 December, the planner reads the Q4 daily rate card by country, slot, date, and buy type. Outside Q4, it continues to use the existing rate card.")
    add_body(doc, "If the daily CPM or CPD rate changes during a campaign, the output is split into separate rows for each continuous date range with the same rate. This makes the booking and partner-facing plan easier to review.")

    add_heading(doc, "When the planner cannot add a slot")
    add_body(doc, "The plan explains why a selected slot could not be added. Common reasons are no forecast inventory in the chosen phase, an unavailable rate, an existing placement conflict, or insufficient remaining budget. For budget cases, the message shows the minimum cost needed and the amount still available.")

    add_heading(doc, "Important trade offs")
    add_body(doc, "Some rules compete with each other. Stronger minimum daily delivery, shorter premium-slot flights, or stricter same-fold exclusions can improve placement quality, but they may reduce budget utilization, make requested splits harder to achieve, or create gaps in delivery. Sales and Operations should agree which outcome takes priority when those rules conflict.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Media Planner Rules Guide", color="7A869A", size=8)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
