from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Media Planner Rules Guide.docx"
NAVY = "17365D"
BLUE = "DCEAF7"
PALE = "F4F7FA"
GRID = "D9D9D9"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def border(cell, color=GRID):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)
        borders.append(element)
    props.append(borders)


def cell_text(cell, text, bold=False, color=None):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.space_before = Pt(3)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(doc, lead, body):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(lead)
    run.bold = True
    paragraph.add_run(body)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Media Planner Prioritization Guide")
# Word's built-in Title style can carry a decorative bottom rule. Suppress it
# explicitly so the title remains plain and readable when opened in Word.
title_borders = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_borders.append(title_bottom)
title._p.get_or_add_pPr().append(title_borders)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("How the planner selects slots and optimizes budget")
run.italic = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(89, 89, 89)

doc.add_paragraph(
    "The planner is designed to create the strongest workable plan, not simply to spend the full budget. "
    "It first protects booking feasibility, then follows the campaign brief, and finally improves the slot mix, daily delivery, and budget use wherever the available inventory allows."
)

doc.add_heading("What the planner prioritizes", level=1)
add_numbered(doc, "A plan that can be booked. ",
             "A slot needs valid inventory, the correct price, and suitable dates. The planner will not create a plan that cannot be delivered or booked.")
add_numbered(doc, "Protection of the approved budget. ",
             "The plan stays within the on-deck budget and avoids concentrating too much spend in one normal paid placement.")
add_numbered(doc, "The agreed campaign brief. ",
             "The planner follows the chosen country, category, phase, marketplace, and campaign objective before looking for extra opportunities.")
add_numbered(doc, "The most relevant slot mix. ",
             "It ranks suitable slots using past performance, category relevance, recent delivery, page type, and confidence in the data. Reach plans favour Homepage visibility; ROAS and CTR plans favour CLP or category placements.")
add_numbered(doc, "Stronger use of the available budget. ",
             "After the above priorities are protected, the planner looks for a balanced mix, useful daily delivery, continuity, and high budget utilization.")

doc.add_heading("Recommendation count", level=1)
doc.add_paragraph(
    "The planner starts with a recommended number of slots for each selected country. It can add more recommendations if the starting set does not use enough of the budget. Users can also add extra slots manually."
)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(2.5)
table.columns[1].width = Inches(4.2)
headers = ("Campaign budget", "Starting recommendation")
for cell, text in zip(table.rows[0].cells, headers):
    shade(cell, NAVY)
    cell_text(cell, text, bold=True, color="FFFFFF")
for budget, recommendation in (
    ("Up to USD 10,000", "6 slots per selected country"),
    ("Above USD 10,000", "10 slots per selected country"),
):
    cells = table.add_row().cells
    cell_text(cells[0], budget)
    cell_text(cells[1], recommendation)

doc.add_heading("CPD rules", level=1)
doc.add_paragraph(
    "CPD means cost per day. The planner checks that the budget can cover at least one valid day before adding a CPD slot.")
add_bullets(doc, [
    "Homepage CPD is recommended only for budgets of USD 15,000 or more. This avoids spreading a premium Homepage placement too thinly across many days.",
    "CLP and category-page CPD can still be recommended below USD 15,000 when there is valid inventory and pricing.",
    "A user can manually choose Homepage CPD below USD 15,000 when there is a commercial reason to do so.",
])

doc.add_heading("Guardrails that keep the plan bookable", level=1)
doc.add_paragraph("These are the limits used while the planner optimizes the plan. They are not relaxed simply to spend more budget.")
add_bullets(doc, [
    "The plan cannot spend more than the on-deck budget.",
    "A normal paid slot cannot take more than 35 percent of the on-deck budget.",
    "A slot must have valid inventory, dates, and pricing for the selected buy type.",
    "The same slot cannot be added twice in the same phase or overlap with itself on the same dates.",
    "The planner avoids duplicate category and zone combinations in the same country and phase.",
    "CPM placements must meet the minimum impression block. CPD placements must be affordable for at least one day.",
])

doc.add_heading("How the planner optimizes within those guardrails", level=1)
doc.add_paragraph(
    "Once a slot is eligible, the planner aims for the closest workable mix. These are targets rather than guarantees because they depend on available inventory, rates, and campaign dates."
)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(4.5)
for cell, text in zip(table.rows[0].cells, ("Target", "What it means")):
    shade(cell, NAVY)
    cell_text(cell, text, bold=True, color="FFFFFF")
for target, meaning in (
    ("Budget use", "Use as much of the budget as practical without breaking booking rules."),
    ("Phase split", "Spread spend across phases in line with the agreed campaign split."),
    ("Marketplace split", "Follow the requested Core and Supermall allocation."),
    ("Campaign mix", "Keep a useful mix of Homepage, CLP, and other relevant pages."),
    ("Daily delivery", "Avoid weak delivery where the available budget and inventory allow a stronger option."),
):
    cells = table.add_row().cells
    cell_text(cells[0], target)
    cell_text(cells[1], meaning)

doc.add_heading("Q4 daily pricing", level=1)
doc.add_paragraph(
    "For October to December, the planner reads the Q4 daily rate card. It uses the correct CPM or CPD rate for each slot, country, and campaign date. If a rate changes during a campaign, the plan shows separate rows for each period with the same rate. This makes the final plan easier to check and share.")

doc.add_heading("When a slot cannot be added", level=1)
doc.add_paragraph(
    "The planner explains the reason in the plan output. Common reasons include unavailable inventory, an invalid rate, a duplicate placement in the same phase, or insufficient remaining budget. Where possible, the message states the minimum amount required and the budget still available.")

doc.add_heading("When priorities conflict", level=1)
doc.add_paragraph(
    "A stricter rule can improve one part of a plan, but it can reduce flexibility elsewhere. For example, a higher minimum daily spend may strengthen a premium placement but make it harder to use the full budget, match every phase or marketplace split, or keep delivery continuous. Before adding a new hard rule, agree which outcome should take priority.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Media Planner Prioritization Guide")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUTPUT)
print(OUTPUT)
